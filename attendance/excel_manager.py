"""
Excel Attendance Manager
-------------------------
Handles writing attendance records to reports/attendance.xlsx.
Columns: Date, Student ID, Name, Roll No, Class, Division, Lecture, Time, Status.

Features:
- Safe migration of legacy 6-column Excel files with automatic backup to reports/attendance_backup.xlsx
- Duplicate prevention per Date + Student ID + Lecture
- Graceful handling of PermissionError when Excel is open in MS Excel:
  saves record to attendance/pending_excel.json and flushes it when Excel is available
- Attendance percentage and statistics calculation (Present + Late count as attended)
"""

import json
import os
import shutil
from datetime import datetime
import pandas as pd
from openpyxl import Workbook, load_workbook

EXCEL_FILE = os.path.join("reports", "attendance.xlsx")
BACKUP_FILE = os.path.join("reports", "attendance_backup.xlsx")
PENDING_FILE = os.path.join("attendance", "pending_excel.json")

COLUMNS = ["Date", "Student ID", "Name", "Roll No", "Class", "Division", "Lecture", "Time", "Status"]


def load_pending_records(pending_file=PENDING_FILE):
    """Load queued records waiting to be written to Excel."""
    if not os.path.exists(pending_file):
        return []
    try:
        with open(pending_file, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, list) else []
    except Exception:
        return []


def save_pending_record(record, pending_file=PENDING_FILE):
    """Queue an attendance record when the Excel file cannot be opened."""
    pending = load_pending_records(pending_file)
    pending.append(record)
    os.makedirs(os.path.dirname(pending_file), exist_ok=True)
    with open(pending_file, "w", encoding="utf-8") as f:
        json.dump(pending, f, indent=2)
    print(f"Record saved to {pending_file}. It will be added to Excel once the file is closed.")


def flush_pending_records(excel_file=EXCEL_FILE, pending_file=PENDING_FILE):
    """Attempt to write any queued pending records to the Excel file."""
    pending = load_pending_records(pending_file)
    if not pending:
        return

    try:
        ensure_excel_exists(excel_file)
        wb = load_workbook(excel_file)
        ws = wb["Attendance"]

        existing_keys = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 7:
                existing_keys.add((str(row[0]), str(row[1]), str(row[6])))

        written = 0
        remaining = []
        for rec in pending:
            key = (str(rec["Date"]), str(rec["Student ID"]), str(rec["Lecture"]))
            if key not in existing_keys:
                ws.append([
                    rec["Date"],
                    rec["Student ID"],
                    rec["Name"],
                    rec.get("Roll No", ""),
                    rec.get("Class", ""),
                    rec.get("Division", ""),
                    rec.get("Lecture", "Old record"),
                    rec["Time"],
                    rec["Status"],
                ])
                existing_keys.add(key)
                written += 1

        wb.save(excel_file)
        wb.close()

        # Remove pending file once saved
        if os.path.exists(pending_file):
            os.remove(pending_file)

        if written > 0:
            print(f"Flushed {written} pending record(s) into {excel_file}")
    except (PermissionError, OSError):
        # Excel is still locked; keep records in pending file
        pass


def migrate_excel_if_needed(excel_file=EXCEL_FILE, backup_file=BACKUP_FILE):
    """
    Checks if the Excel file uses the old 6-column schema:
    [Date, Student ID, Name, Roll No, Time, Status]
    If so, backs it up and migrates it to the 9-column schema:
    [Date, Student ID, Name, Roll No, Class, Division, Lecture, Time, Status]
    """
    if not os.path.exists(excel_file):
        return

    try:
        wb = load_workbook(excel_file)
        ws = wb["Attendance"]
        header = [cell.value for cell in ws[1] if cell.value is not None]

        # If it's already the new 9-column layout, nothing to do
        if header == COLUMNS:
            wb.close()
            return

        # Old 6-column schema detected
        if len(header) == 6 and header[:4] == ["Date", "Student ID", "Name", "Roll No"]:
            print(f"Migrating {excel_file} from 6-column schema to 9-column schema...")
            # Make a safe backup copy
            os.makedirs(os.path.dirname(backup_file), exist_ok=True)
            shutil.copyfile(excel_file, backup_file)
            print(f"Backup saved to: {backup_file}")

            rows_to_keep = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                date_val = str(row[0]) if row[0] is not None else ""
                sid_val = str(row[1]) if row[1] is not None else ""
                name_val = str(row[2]) if row[2] is not None else ""
                roll_val = str(row[3]) if row[3] is not None else ""
                time_val = str(row[4]) if row[4] is not None else ""
                status_val = str(row[5]) if row[5] is not None else "Present"

                # New columns: Date, Student ID, Name, Roll No, Class, Division, Lecture, Time, Status
                rows_to_keep.append([
                    date_val,
                    sid_val,
                    name_val,
                    roll_val,
                    "",           # Class (empty for old records)
                    "",           # Division (empty for old records)
                    "Old record", # Lecture default
                    time_val,
                    status_val
                ])

            wb.close()

            # Create fresh workbook with new header and converted rows
            new_wb = Workbook()
            new_ws = new_wb.active
            new_ws.title = "Attendance"
            new_ws.append(COLUMNS)
            for r in rows_to_keep:
                new_ws.append(r)
            new_wb.save(excel_file)
            new_wb.close()
            print(f"Migration completed successfully. {len(rows_to_keep)} rows preserved.")
        else:
            wb.close()
    except (PermissionError, OSError):
        print(f"Please close {excel_file} and try again.")
    except Exception as e:
        print(f"Warning during Excel check: {e}")


def ensure_excel_exists(excel_file=EXCEL_FILE):
    """Create the Excel file with headers if it doesn't already exist, or migrate old schema."""
    os.makedirs(os.path.dirname(excel_file), exist_ok=True)

    if not os.path.exists(excel_file):
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(COLUMNS)
        wb.save(excel_file)
        wb.close()
        print(f"Created new Excel file: {excel_file}")
    else:
        migrate_excel_if_needed(excel_file)


def add_attendance_record(student_id, name, roll_no="", student_class="", division="",
                          lecture="Old record", status_date=None, time_str=None,
                          status="Present", excel_file=EXCEL_FILE):
    """
    Append a new attendance row to the Excel file.
    Duplicate check: avoids duplicate (Date + Student ID + Lecture).
    If Excel is open, saves record in pending_excel.json without crashing.
    """
    if status_date is None:
        status_date = datetime.now().strftime("%Y-%m-%d")
    if time_str is None:
        time_str = datetime.now().strftime("%H:%M:%S")

    record_dict = {
        "Date": str(status_date),
        "Student ID": str(student_id),
        "Name": str(name),
        "Roll No": str(roll_no) if roll_no is not None else "",
        "Class": str(student_class) if student_class is not None else "",
        "Division": str(division) if division is not None else "",
        "Lecture": str(lecture) if lecture is not None else "Old record",
        "Time": str(time_str),
        "Status": str(status),
    }

    try:
        ensure_excel_exists(excel_file)
        flush_pending_records(excel_file)

        wb = load_workbook(excel_file)
        ws = wb["Attendance"]

        # Duplicate check: Date (col 1), Student ID (col 2), Lecture (col 7)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 7:
                row_date = str(row[0]) if row[0] is not None else ""
                row_id = str(row[1]) if row[1] is not None else ""
                row_lecture = str(row[6]) if row[6] is not None else ""

                if row_date == str(status_date) and row_id == str(student_id) and row_lecture == str(lecture):
                    wb.close()
                    return False  # Already exists

        new_row = [
            record_dict["Date"],
            record_dict["Student ID"],
            record_dict["Name"],
            record_dict["Roll No"],
            record_dict["Class"],
            record_dict["Division"],
            record_dict["Lecture"],
            record_dict["Time"],
            record_dict["Status"],
        ]
        ws.append(new_row)
        wb.save(excel_file)
        wb.close()
        return True

    except (PermissionError, OSError):
        print(f"Please close {os.path.basename(excel_file)} and try again.")
        save_pending_record(record_dict)
        return False
    except Exception as e:
        print(f"Error writing to Excel: {e}")
        save_pending_record(record_dict)
        return False


def read_attendance(excel_file=EXCEL_FILE):
    """Read all attendance records as a Pandas DataFrame."""
    ensure_excel_exists(excel_file)
    flush_pending_records(excel_file)
    try:
        df = pd.read_excel(excel_file, sheet_name="Attendance")
        return df
    except Exception as e:
        print(f"Error reading attendance records: {e}")
        return pd.DataFrame(columns=COLUMNS)


def get_attendance_percentage(student_id, total_lectures, excel_file=EXCEL_FILE):
    """
    Calculate a student's attendance percentage.
    Both 'Present' and 'Late' count as attended.
    Returns: (percentage, present_and_late_count, absent_count)
    """
    df = read_attendance(excel_file)
    if df.empty:
        return 0.0, 0, total_lectures

    # Filter for this student where Status is 'Present' or 'Late'
    mask = (df["Student ID"].astype(str) == str(student_id)) & (df["Status"].isin(["Present", "Late"]))
    attended_count = len(df[mask])

    if total_lectures <= 0:
        return 0.0, attended_count, 0

    percentage = (attended_count / total_lectures) * 100
    absent_count = max(0, total_lectures - attended_count)
    return round(percentage, 2), attended_count, absent_count


def filter_by_date(date_str, excel_file=EXCEL_FILE):
    """Return all attendance records for a specific date (YYYY-MM-DD)."""
    df = read_attendance(excel_file)
    if df.empty:
        return df
    return df[df["Date"].astype(str) == str(date_str)]


def filter_by_student(student_id, excel_file=EXCEL_FILE):
    """Return all attendance records for a specific student."""
    df = read_attendance(excel_file)
    if df.empty:
        return df
    return df[df["Student ID"].astype(str) == str(student_id)]


def filter_by_lecture(lecture_name, excel_file=EXCEL_FILE):
    """Return all attendance records for a specific lecture."""
    df = read_attendance(excel_file)
    if df.empty:
        return df
    return df[df["Lecture"].astype(str) == str(lecture_name)]


if __name__ == "__main__":
    print("Testing excel_manager migration and read...")
    ensure_excel_exists()
    df = read_attendance()
    print("Current columns in Excel:", list(df.columns))
    print(df.head())